"""Bounded public-data reproduction spike for the published gastric NRRS.

This is not an exact reconstruction of the anchor paper.  The paper does not
specify its probe-collapsing rule, z-score convention, or high/low cutoff.  The
main analysis here therefore locks an explicit, auditable contract:

* use the GEO RMA series matrix as distributed;
* map probes with the dated NCBI GPL570 annotation;
* retain probes assigned unambiguously to one target gene;
* average all retained probes for each gene;
* z-score each gene within GSE62254 with the sample standard deviation;
* apply the published eight coefficients in the published direction; and
* split at the cohort median (ties are assigned deterministically by GEO ID).

A highest-variance-probe calculation is emitted as a mapping sensitivity check.
GSE62254 participated in outcome-guided feature screening in the anchor study,
so the resulting association is a reproduction spike, not independent external
validation.
"""

from __future__ import annotations

import argparse
import csv
import gzip
import hashlib
import importlib.metadata
import json
import platform
import sys
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable
from urllib.request import Request, urlopen

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from lifelines import CoxPHFitter, KaplanMeierFitter
from lifelines.statistics import logrank_test
from lifelines.utils import concordance_index
from openpyxl import load_workbook


COEFFICIENTS = {
    "AGT": 0.040,
    "EPHB3": -0.040,
    "GNAI1": 0.018,
    "LPAR2": -0.102,
    "LRRC4C": 0.014,
    "NPY1R": 0.072,
    "NRP1": 0.175,
    "SEMA6A": 0.006,
}
TARGET_GENES = tuple(COEFFICIENTS)
CLINICAL_SHEET = "ACRG"


@dataclass(frozen=True)
class Resource:
    resource_id: str
    url: str
    local_name: str
    expected_bytes: int
    sha256: str


def sha256_file(path: Path, chunk_size: int = 1024 * 1024) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        while chunk := handle.read(chunk_size):
            digest.update(chunk)
    return digest.hexdigest().upper()


def metadata_path(path: Path, base: Path | None = None) -> str:
    """Return an auditable path without exposing an absolute local directory."""
    resolved = path.resolve()
    root = (base or Path.cwd()).resolve()
    try:
        return resolved.relative_to(root).as_posix()
    except ValueError:
        return resolved.name


def load_resource_manifest(path: Path) -> list[Resource]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle, delimiter="\t"))
    resources = []
    for row in rows:
        resources.append(
            Resource(
                resource_id=row["resource_id"],
                url=row["url"],
                local_name=row["local_name"],
                expected_bytes=int(row["bytes"]),
                sha256=row["sha256"].upper(),
            )
        )
    if not resources:
        raise ValueError("resource manifest is empty")
    return resources


def verify_resource(path: Path, resource: Resource) -> dict[str, object]:
    observed_bytes = path.stat().st_size
    observed_sha256 = sha256_file(path)
    if observed_bytes != resource.expected_bytes:
        raise ValueError(
            f"{resource.resource_id}: expected {resource.expected_bytes} bytes, "
            f"observed {observed_bytes}"
        )
    if observed_sha256 != resource.sha256:
        raise ValueError(
            f"{resource.resource_id}: checksum mismatch; public source may have changed"
        )
    return {
        "resource_id": resource.resource_id,
        "path": metadata_path(path),
        "bytes": observed_bytes,
        "sha256": observed_sha256,
        "verified": True,
    }


def download_resources(manifest_path: Path, output_dir: Path) -> list[dict[str, object]]:
    output_dir.mkdir(parents=True, exist_ok=True)
    reports = []
    for resource in load_resource_manifest(manifest_path):
        destination = output_dir / resource.local_name
        if not destination.exists():
            temporary = destination.with_suffix(destination.suffix + ".part")
            request = Request(resource.url, headers={"User-Agent": "Paper2Paper/0.4"})
            with urlopen(request, timeout=120) as response, temporary.open("wb") as target:
                while chunk := response.read(1024 * 1024):
                    target.write(chunk)
            temporary.replace(destination)
        reports.append(verify_resource(destination, resource))
    return reports


def parse_geo_matrix_metadata(matrix_path: Path) -> pd.DataFrame:
    sample_titles: list[str] = []
    sample_ids: list[str] = []
    patient_ids: list[str] = []
    with gzip.open(
        matrix_path, "rt", encoding="utf-8", errors="replace", newline=""
    ) as handle:
        for raw_line in handle:
            line = raw_line.rstrip("\r\n")
            if line == "!series_matrix_table_begin":
                break
            fields = next(csv.reader([line], delimiter="\t", quotechar='"'))
            if not fields:
                continue
            if fields[0] == "!Sample_title":
                sample_titles = fields[1:]
            elif fields[0] == "!Sample_geo_accession":
                sample_ids = fields[1:]
            elif fields[0] == "!Sample_characteristics_ch1" and fields[1:]:
                if fields[1].startswith("patient:"):
                    patient_ids = [value.split(":", 1)[1].strip() for value in fields[1:]]
    lengths = {len(sample_titles), len(sample_ids), len(patient_ids)}
    if lengths != {300}:
        raise ValueError(
            "GSE62254 metadata contract failed: expected 300 titles, GEO IDs and patient IDs"
        )
    metadata = pd.DataFrame(
        {
            "gsm_id": sample_ids,
            "sample_title": sample_titles,
            "patient_id": patient_ids,
        }
    )
    if metadata["gsm_id"].duplicated().any() or metadata["patient_id"].duplicated().any():
        raise ValueError("GSE62254 matrix contains duplicate sample or patient identifiers")
    return metadata


def parse_probe_map(annotation_path: Path) -> pd.DataFrame:
    rows: list[dict[str, str]] = []
    with gzip.open(
        annotation_path, "rt", encoding="utf-8", errors="replace", newline=""
    ) as handle:
        reader: Iterable[list[str]] | None = None
        gene_index = -1
        for line in handle:
            if line.startswith("ID\t"):
                header = next(csv.reader([line], delimiter="\t"))
                gene_index = header.index("Gene symbol")
                reader = csv.reader(handle, delimiter="\t")
                break
        if reader is None:
            raise ValueError("GPL570 annotation header was not found")
        for fields in reader:
            if len(fields) <= gene_index:
                continue
            symbols = [value.strip() for value in fields[gene_index].split("///")]
            symbols = [value for value in symbols if value and value != "---"]
            matched = [value for value in symbols if value in COEFFICIENTS]
            if len(symbols) == 1 and len(matched) == 1:
                rows.append({"probe_id": fields[0], "gene": matched[0]})
    mapping = pd.DataFrame(rows).sort_values(["gene", "probe_id"]).reset_index(drop=True)
    observed = set(mapping["gene"])
    missing = sorted(set(TARGET_GENES) - observed)
    if missing:
        raise ValueError(f"target genes without unambiguous GPL570 probes: {missing}")
    if mapping["probe_id"].duplicated().any():
        raise ValueError("probe was assigned to more than one target gene")
    return mapping


def read_target_probe_expression(
    matrix_path: Path, sample_ids: list[str], probe_map: pd.DataFrame
) -> pd.DataFrame:
    wanted = set(probe_map["probe_id"])
    rows: list[list[object]] = []
    table_started = False
    with gzip.open(
        matrix_path, "rt", encoding="utf-8", errors="replace", newline=""
    ) as handle:
        for raw_line in handle:
            line = raw_line.rstrip("\r\n")
            if line == "!series_matrix_table_begin":
                table_started = True
                continue
            if line == "!series_matrix_table_end":
                break
            if not table_started:
                continue
            fields = next(csv.reader([line], delimiter="\t", quotechar='"'))
            if fields[0] == "ID_REF":
                if fields[1:] != sample_ids:
                    raise ValueError("matrix expression columns do not match metadata GEO IDs")
                continue
            if fields[0] in wanted:
                if len(fields) != len(sample_ids) + 1:
                    raise ValueError(f"malformed expression row for {fields[0]}")
                rows.append([fields[0], *[float(value) for value in fields[1:]]])
    expression = pd.DataFrame(rows, columns=["probe_id", *sample_ids]).set_index("probe_id")
    missing = sorted(wanted - set(expression.index))
    if missing:
        raise ValueError(f"target probes missing from matrix: {missing}")
    if not np.isfinite(expression.to_numpy(dtype=float)).all():
        raise ValueError("target expression contains non-finite values")
    return expression


def load_acrg_clinical(clinical_path: Path) -> pd.DataFrame:
    workbook = load_workbook(clinical_path, read_only=True, data_only=True)
    if CLINICAL_SHEET not in workbook.sheetnames:
        raise ValueError(f"clinical workbook lacks {CLINICAL_SHEET!r} sheet")
    sheet = workbook[CLINICAL_SHEET]
    header = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    required = {"GEO_ID", "Death", "OS.m"}
    if not required.issubset(header):
        raise ValueError(f"clinical sheet lacks fields: {sorted(required - set(header))}")
    positions = {name: header.index(name) for name in required}
    rows = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        gsm_id = values[positions["GEO_ID"]]
        if gsm_id in (None, ""):
            continue
        rows.append(
            {
                "gsm_id": str(gsm_id).strip(),
                "death": int(values[positions["Death"]]),
                "os_months": float(values[positions["OS.m"]]),
            }
        )
    clinical = pd.DataFrame(rows)
    if len(clinical) != 300 or clinical["gsm_id"].nunique() != 300:
        raise ValueError("ACRG clinical sheet must contain 300 unique GEO IDs")
    if not set(clinical["death"]).issubset({0, 1}):
        raise ValueError("Death must be encoded as 0/1")
    if clinical[["death", "os_months"]].isna().any().any():
        raise ValueError("ACRG survival fields contain missing values")
    if (clinical["os_months"] <= 0).any():
        raise ValueError("ACRG OS time must be positive")
    return clinical


def aggregate_gene_expression(
    probe_expression: pd.DataFrame, probe_map: pd.DataFrame, method: str
) -> tuple[pd.DataFrame, pd.DataFrame]:
    mapping = probe_map.copy()
    variances = probe_expression.var(axis=1, ddof=1)
    mapping["sample_variance"] = mapping["probe_id"].map(variances)
    mapping["aggregation"] = method
    selected_frames = []
    for gene in TARGET_GENES:
        gene_map = mapping[mapping["gene"] == gene].copy()
        if method == "mean_probes":
            chosen = gene_map["probe_id"].tolist()
            values = probe_expression.loc[chosen].mean(axis=0)
            gene_map["selected"] = True
        elif method == "max_variance_probe":
            gene_map = gene_map.sort_values(
                ["sample_variance", "probe_id"], ascending=[False, True]
            )
            chosen = [gene_map.iloc[0]["probe_id"]]
            values = probe_expression.loc[chosen[0]]
            gene_map["selected"] = gene_map["probe_id"].isin(chosen)
        else:
            raise ValueError(f"unknown aggregation method: {method}")
        selected_frames.append(values.rename(gene))
        mapping.loc[gene_map.index, "selected"] = gene_map["selected"]
    genes_by_sample = pd.concat(selected_frames, axis=1)
    genes_by_sample.index.name = "gsm_id"
    return genes_by_sample, mapping


def compute_nrrs(gene_expression: pd.DataFrame) -> tuple[pd.DataFrame, pd.Series]:
    missing = [gene for gene in TARGET_GENES if gene not in gene_expression.columns]
    if missing:
        raise ValueError(f"gene expression lacks NRRS genes: {missing}")
    means = gene_expression.loc[:, TARGET_GENES].mean(axis=0)
    standard_deviations = gene_expression.loc[:, TARGET_GENES].std(axis=0, ddof=1)
    if (standard_deviations <= 0).any():
        raise ValueError("cannot z-score a zero-variance NRRS gene")
    zscores = (gene_expression.loc[:, TARGET_GENES] - means) / standard_deviations
    coefficients = pd.Series(COEFFICIENTS)
    score = zscores.mul(coefficients, axis=1).sum(axis=1).rename("nrrs")
    return zscores, score


def median_groups(score: pd.Series) -> tuple[pd.Series, float]:
    cutoff = float(score.median())
    groups = pd.Series("low", index=score.index, name="risk_group")
    groups.loc[score > cutoff] = "high"
    ties = score.index[score == cutoff].tolist()
    if ties:
        target_high = len(score) // 2
        need = target_high - int((groups == "high").sum())
        for sample_id in sorted(ties)[: max(need, 0)]:
            groups.loc[sample_id] = "high"
    if abs(int((groups == "high").sum()) - int((groups == "low").sum())) > 1:
        raise ValueError("median grouping did not produce balanced groups")
    return groups, cutoff


def survival_statistics(frame: pd.DataFrame, score_column: str) -> dict[str, float]:
    analysis = frame[["os_months", "death", score_column, "risk_group"]].copy()
    score_sd = analysis[score_column].std(ddof=1)
    analysis["score_per_sd"] = (
        analysis[score_column] - analysis[score_column].mean()
    ) / score_sd
    analysis["high"] = (analysis["risk_group"] == "high").astype(int)

    continuous = CoxPHFitter()
    continuous.fit(
        analysis[["os_months", "death", "score_per_sd"]],
        duration_col="os_months",
        event_col="death",
    )
    binary = CoxPHFitter()
    binary.fit(
        analysis[["os_months", "death", "high"]],
        duration_col="os_months",
        event_col="death",
    )
    high = analysis[analysis["high"] == 1]
    low = analysis[analysis["high"] == 0]
    logrank = logrank_test(
        high["os_months"],
        low["os_months"],
        event_observed_A=high["death"],
        event_observed_B=low["death"],
    )
    continuous_row = continuous.summary.loc["score_per_sd"]
    binary_row = binary.summary.loc["high"]
    return {
        "subjects": float(len(analysis)),
        "events": float(analysis["death"].sum()),
        "high_n": float(len(high)),
        "low_n": float(len(low)),
        "continuous_hr_per_sd": float(continuous_row["exp(coef)"]),
        "continuous_ci_lower": float(continuous_row["exp(coef) lower 95%"]),
        "continuous_ci_upper": float(continuous_row["exp(coef) upper 95%"]),
        "continuous_p": float(continuous_row["p"]),
        "high_vs_low_hr": float(binary_row["exp(coef)"]),
        "high_vs_low_ci_lower": float(binary_row["exp(coef) lower 95%"]),
        "high_vs_low_ci_upper": float(binary_row["exp(coef) upper 95%"]),
        "high_vs_low_p": float(binary_row["p"]),
        "logrank_chi2": float(logrank.test_statistic),
        "logrank_p": float(logrank.p_value),
        "c_index": float(
            concordance_index(
                analysis["os_months"], -analysis[score_column], analysis["death"]
            )
        ),
    }


def plot_kaplan_meier(frame: pd.DataFrame, output: Path, logrank_p: float) -> None:
    figure, axis = plt.subplots(figsize=(6.2, 4.8))
    colors = {"high": "#C44E52", "low": "#4C72B0"}
    for group in ("high", "low"):
        subset = frame[frame["risk_group"] == group]
        estimator = KaplanMeierFitter(label=f"{group.title()} (n={len(subset)})")
        estimator.fit(subset["os_months"], subset["death"])
        estimator.plot_survival_function(ax=axis, ci_show=True, color=colors[group])
    axis.set_xlabel("Overall survival (months)")
    axis.set_ylabel("Survival probability")
    axis.set_ylim(0, 1.02)
    axis.set_title("GSE62254 reconstructed NRRS median split")
    axis.text(
        0.98,
        0.05,
        f"Log-rank p={logrank_p:.3g}",
        transform=axis.transAxes,
        horizontalalignment="right",
    )
    figure.tight_layout()
    output.parent.mkdir(parents=True, exist_ok=True)
    figure.savefig(output, dpi=180)
    plt.close(figure)


def write_tsv(frame: pd.DataFrame, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    frame.to_csv(path, sep="\t", index=False, lineterminator="\n")


def verify_expected_results(
    observed: dict[str, object], expected_path: Path
) -> dict[str, object]:
    expected = json.loads(expected_path.read_text(encoding="utf-8"))
    checked = 0

    def compare(expected_value: object, observed_value: object, key: str) -> None:
        nonlocal checked
        if isinstance(expected_value, dict):
            if not isinstance(observed_value, dict):
                raise ValueError(f"expected-result type mismatch at {key}")
            for child_key, child_expected in expected_value.items():
                if child_key not in observed_value:
                    raise ValueError(f"expected-result key missing: {key}.{child_key}")
                compare(
                    child_expected,
                    observed_value[child_key],
                    f"{key}.{child_key}",
                )
            return
        checked += 1
        if isinstance(expected_value, bool):
            if observed_value is not expected_value:
                raise ValueError(f"expected-result mismatch at {key}")
            return
        if isinstance(expected_value, (int, float)):
            if not np.isclose(
                float(observed_value), float(expected_value), rtol=1e-7, atol=1e-10
            ):
                raise ValueError(
                    f"expected-result mismatch at {key}: "
                    f"expected {expected_value}, observed {observed_value}"
                )
            return
        if observed_value != expected_value:
            raise ValueError(f"expected-result mismatch at {key}")

    compare(expected, observed, "results")
    return {
        "path": metadata_path(expected_path),
        "sha256": sha256_file(expected_path),
        "metrics_checked": checked,
        "verified": True,
    }


def analyze(
    manifest_path: Path,
    data_dir: Path,
    output_dir: Path,
) -> dict[str, object]:
    resources = {item.resource_id: item for item in load_resource_manifest(manifest_path)}
    required_ids = {"GSE62254_MATRIX", "GPL570_ANNOTATION", "ACRG_CLINICAL"}
    if set(resources) != required_ids:
        raise ValueError(f"manifest resource IDs must be {sorted(required_ids)}")
    verified = {
        key: verify_resource(data_dir / resources[key].local_name, resources[key])
        for key in sorted(required_ids)
    }
    matrix_path = data_dir / resources["GSE62254_MATRIX"].local_name
    annotation_path = data_dir / resources["GPL570_ANNOTATION"].local_name
    clinical_path = data_dir / resources["ACRG_CLINICAL"].local_name

    metadata = parse_geo_matrix_metadata(matrix_path)
    probe_map = parse_probe_map(annotation_path)
    probe_expression = read_target_probe_expression(
        matrix_path, metadata["gsm_id"].tolist(), probe_map
    )
    clinical = load_acrg_clinical(clinical_path)
    if set(metadata["gsm_id"]) != set(clinical["gsm_id"]):
        raise ValueError("GEO matrix and clinical supplement GEO IDs do not match 1:1")

    outputs: dict[str, object] = {}
    score_frames: dict[str, pd.DataFrame] = {}
    mapping_frames = []
    for method in ("mean_probes", "max_variance_probe"):
        genes, mapping = aggregate_gene_expression(probe_expression, probe_map, method)
        zscores, score = compute_nrrs(genes)
        groups, cutoff = median_groups(score)
        frame = (
            metadata.set_index("gsm_id")
            .join(clinical.set_index("gsm_id"), how="inner")
            .join(genes.add_prefix("expression_"))
            .join(zscores.add_prefix("z_"))
            .join(score)
            .join(groups)
            .reset_index()
        )
        if len(frame) != 300:
            raise ValueError("joined source table does not contain 300 patients")
        statistics = survival_statistics(frame, "nrrs")
        statistics["median_cutoff"] = cutoff
        statistics["aggregation"] = method
        statistics["risk_direction_matches_anchor"] = statistics["high_vs_low_hr"] > 1
        outputs[method] = statistics
        mapping_frames.append(mapping)
        score_frames[method] = frame

    main = score_frames["mean_probes"].copy()
    sensitivity = score_frames["max_variance_probe"].set_index("gsm_id")["nrrs"]
    main["nrrs_max_variance_probe"] = main["gsm_id"].map(sensitivity)
    score_correlation = float(
        main[["nrrs", "nrrs_max_variance_probe"]].corr(method="spearman").iloc[0, 1]
    )
    outputs["mapping_sensitivity"] = {
        "spearman_score_correlation": score_correlation,
        "high_low_agreement": float(
            (
                main.set_index("gsm_id")["risk_group"]
                == score_frames["max_variance_probe"].set_index("gsm_id")["risk_group"]
            ).mean()
        ),
    }

    output_dir.mkdir(parents=True, exist_ok=True)
    score_columns = [
        "gsm_id",
        "patient_id",
        "os_months",
        "death",
        *[f"expression_{gene}" for gene in TARGET_GENES],
        *[f"z_{gene}" for gene in TARGET_GENES],
        "nrrs",
        "risk_group",
        "nrrs_max_variance_probe",
    ]
    write_tsv(main[score_columns], output_dir / "GSE62254_nrrs_source_table.tsv")
    mapping_output = pd.concat(mapping_frames, ignore_index=True)
    mapping_output["selected"] = mapping_output["selected"].astype(bool)
    write_tsv(mapping_output, output_dir / "GSE62254_probe_mapping.tsv")
    summary_rows = []
    for method in ("mean_probes", "max_variance_probe"):
        summary_rows.append(
            {"analysis": method, **outputs[method]}  # type: ignore[arg-type]
        )
    summary_frame = pd.DataFrame(summary_rows)
    write_tsv(summary_frame, output_dir / "GSE62254_nrrs_summary.tsv")
    expected_path = manifest_path.with_name("gse62254_expected_results.json")
    if not expected_path.exists():
        raise ValueError(f"expected-result contract is missing: {expected_path}")
    expected_result_check = verify_expected_results(outputs, expected_path)
    plot_kaplan_meier(
        main,
        output_dir / "GSE62254_nrrs_kaplan_meier.png",
        float(outputs["mean_probes"]["logrank_p"]),  # type: ignore[index]
    )

    run_metadata = {
        "run_at_utc": datetime.now(timezone.utc).isoformat(),
        "code": {
            "path": metadata_path(Path(__file__)),
            "sha256": sha256_file(Path(__file__)),
        },
        "python": sys.version.split()[0],
        "platform": platform.platform(),
        "packages": {
            "numpy": np.__version__,
            "pandas": pd.__version__,
            "matplotlib": matplotlib.__version__,
            "lifelines": importlib.metadata.version("lifelines"),
            "openpyxl": importlib.metadata.version("openpyxl"),
            "scipy": importlib.metadata.version("scipy"),
        },
        "resources": verified,
        "model": {
            "formula": " + ".join(
                f"{coefficient:.3f}*z({gene})"
                for gene, coefficient in COEFFICIENTS.items()
            ),
            "main_probe_rule": "mean of all unambiguous official GPL570 probes",
            "normalization": "within-cohort gene z-score; sample SD (ddof=1)",
            "cutoff": "cohort median with deterministic GEO-ID tie handling",
            "missing_gene_policy": "fail",
        },
        "identity_checks": {
            "matrix_samples": len(metadata),
            "clinical_samples": len(clinical),
            "exact_geo_id_set_match": True,
            "target_genes": len(TARGET_GENES),
            "target_probes": len(probe_map),
        },
        "results": outputs,
        "expected_result_check": expected_result_check,
        "claim_boundary": (
            "Reproduction spike only; GSE62254 participated in anchor feature "
            "screening and is not untouched external validation."
        ),
    }
    (output_dir / "GSE62254_run_metadata.json").write_text(
        json.dumps(run_metadata, indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )
    return run_metadata


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    subparsers = parser.add_subparsers(dest="command", required=True)
    download_parser = subparsers.add_parser("download")
    download_parser.add_argument("--manifest", type=Path, required=True)
    download_parser.add_argument("--output-dir", type=Path, required=True)
    analyze_parser = subparsers.add_parser("analyze")
    analyze_parser.add_argument("--manifest", type=Path, required=True)
    analyze_parser.add_argument("--data-dir", type=Path, required=True)
    analyze_parser.add_argument("--output-dir", type=Path, required=True)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    if args.command == "download":
        print(json.dumps(download_resources(args.manifest, args.output_dir), indent=2))
    elif args.command == "analyze":
        result = analyze(args.manifest, args.data_dir, args.output_dir)
        print(json.dumps(result["results"], indent=2))


if __name__ == "__main__":
    main()
